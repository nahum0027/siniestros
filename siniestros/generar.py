"""
Llena los 3 formatos de siniestro a partir de un diccionario `datos`.

Uso:
    from siniestros.generar import generar_todo
    archivos = generar_todo(datos, carpeta_salida)

Las firmas y el croquis llegan como rutas a PNG (o None).
"""
import re
from pathlib import Path
from datetime import date

from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill, Font, Alignment

PLANTILLAS = Path(__file__).resolve().parent / "plantillas"

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def _img(doc, ruta, ancho_mm=None, alto_mm=None):
    if not ruta:
        return ""
    if alto_mm:
        return InlineImage(doc, str(ruta), height=Mm(alto_mm))
    return InlineImage(doc, str(ruta), width=Mm(ancho_mm))


# ---------------------------------------------------------------- WORD
def llenar_convenio(datos, salida):
    doc = DocxTemplate(PLANTILLAS / "convenio.docx")
    f = datos["fecha"]  # date
    ctx = {
        "dia": f.day, "mes": MESES[f.month - 1], "anio": f.year,
        "narrativa": datos.get("narrativa", ""),
        "forma_reparacion": datos.get("forma_reparacion", ""),
        "agente_nombre": datos.get("agente_nombre", ""),
        "firma_a": _img(doc, datos.get("firma_conv_a"), 40),
        "firma_b": _img(doc, datos.get("firma_conv_b"), 40),
        "firma_agente": _img(doc, datos.get("firma_agente"), 40),
    }
    for k in ("a", "b"):
        p = datos[k]
        ctx.update({f"{k}_{c}": p.get(c, "") for c in
                    ("nombre", "domicilio", "telefono", "marca", "modelo", "placas", "identificacion")})
    doc.render(ctx)
    ruta = Path(salida) / "convenio.docx"
    doc.save(ruta)
    return ruta


def llenar_entrevista(datos, salida):
    doc = DocxTemplate(PLANTILLAS / "entrevista.docx")
    a = datos["a"]
    ctx = {
        "agente_nombre": datos.get("agente_nombre", ""),
        "fecha": datos["fecha"].strftime("%d/%m/%Y"),
        "hora": datos.get("hora", ""),
        "lugar": datos.get("lugar", ""),
        "economico": datos.get("economico", ""),
        "placa": a.get("placas", ""),
        "ruta": datos.get("ruta", ""),
        "no_empleado": a.get("no_empleado", ""),
        "tel_ejecutivo": a.get("telefono", ""),
        "reporte_aseguradora": datos.get("reporte_aseguradora", ""),
        "narrativa": datos.get("narrativa", ""),
        "seguro_activo": datos.get("seguro_activo", ""),
        "pases_medicos": datos.get("pases_medicos", ""),
        "causal": datos.get("causal", ""),
        "seguro_garantiza": datos.get("seguro_garantiza", ""),
        "ejecutivo_nombre": a.get("nombre", ""),
        "croquis": _img(doc, datos.get("croquis"), alto_mm=62),
        "firma_ejecutivo": _img(doc, datos.get("firma_ejecutivo"), 45),
    }
    doc.render(ctx)
    ruta = Path(salida) / "entrevista.docx"
    doc.save(ruta)
    return ruta


# ---------------------------------------------------------------- EXCEL
# Celdas donde el valor va solo (celda vacía en el formato)
VALOR = {
    "dia": "B7", "mes": "E7", "anio": "H7", "hh": "Q7", "mm": "T7", "ruta_eco": "W7",
    "a_nombre": "A12", "a_edad": "P12", "a_telefono": "X12",
    "a_tarjeton": "Z14", "a_placas": "A17", "a_modelo": "P17", "a_garantiza": "Y17",
    "a_trayectoria": "G18",
    "b_nombre": "A21", "b_edad": "P21", "b_telefono": "X21", "b_domicilio": "D22",
    "b_placas": "A24", "b_entidad": "E24", "b_marca": "I24", "b_tipo_servicio": "L24",
    "b_modelo": "P24", "b_seguro": "T24", "b_garantiza": "Y24", "b_trayectoria": "G25",
    "causal": "D38", "colision": "D39", "medio_prueba": "D41",
}
# Celdas que ya traen una etiqueta chiquita: el valor se agrega después de la etiqueta
ETIQUETA = {
    "reporte_aseguradora": "R4", "lugar": "A8", "colonia": "R8",
    "a_licencia": "N13", "a_vigencia": "S13", "a_infraccion": "Z13", "a_garantia": "AC13",
    "a_remolcado_por": "A19", "a_trasladado_a": "K19", "a_disposicion": "V19",
    "b_remolcado_por": "A26", "b_trasladado_a": "K26", "b_disposicion": "V26",
    "nota": "Q38",
    "ejecutivo_nombre": "A53", "agente_nombre": "L53", "perito_nombre": "V53",
}
DIAS_SEMANA = ["J7", "K7", "L7", "M7", "N7", "O7", "P7"]  # L M M J V S D
COMPETENCIA = {"MUNICIPAL": "F9", "ESTATAL": "K9", "GN": "O9", "OTROS": "S9"}
SEXO = {"a": {"M": "N12", "F": "O12"}, "b": {"M": "N21", "F": "O21"}}
NACIMIENTO = {"a": ["R12", "S12", "T12", "U12", "V12", "W12"],
              "b": ["R21", "S21", "T21", "U21", "V21", "W21"]}
LESIONADO = {"a": "AC11", "b": "AC20"}
FALLECIDO = {"a": "AC12", "b": "AC21"}
RESPONSABLE = {True: "D40", False: "E40"}
LESIONADOS_FILAS = range(29, 37)  # 8 renglones
LESIONADOS_COLS = ("nombre", "A"), ("direccion", "J"), ("telefono", "S"), ("pase_medico", "W"), ("trasladado_a", "AA")
ANEXOS = range(43, 49)  # 6 leyes/reglamentos
FIRMAS = {"firma_ejecutivo": "A52", "firma_agente": "L52", "firma_perito": "V52"}

MARCA = PatternFill("solid", fgColor="BFBFBF")


def _escribir(ws, coord, valor, sz=9):
    """Escribe el valor con letra legible (varias celdas del formato traen letra de 5-6 pts)."""
    c = ws[coord]
    c.value = valor
    c.font = Font(name=c.font.name, sz=sz, bold=c.font.b)


# Celdas angostas: la etiqueta queda arriba y el valor en un segundo renglón
DOS_RENGLONES = {"N13", "S13", "Z13", "AC13"}
# Celdas de texto libre que pueden ser largas -> caracteres aprox. por renglón
CRECE = {"Q38": 62}


def _marcar(ws, coord):
    """Sombrea y pone en negritas una opción (día, sexo, SI/NO, competencia)."""
    c = ws[coord]
    c.fill = MARCA
    c.font = Font(name=c.font.name, sz=c.font.sz, bold=True)


def _con_etiqueta(ws, coord, valor):
    c = ws[coord]
    etiqueta = str(c.value or "").strip()
    f = c.font
    chica = InlineFont(rFont=f.name, sz=f.sz, b=f.b)
    grande = InlineFont(rFont=f.name, sz=9, b=True)
    sep = "\n" if coord in DOS_RENGLONES else "  "
    c.value = CellRichText(TextBlock(chica, etiqueta + sep), TextBlock(grande, str(valor)))
    if coord in CRECE:
        # texto libre largo (p. ej. la NOTA): salto de línea y el renglón crece según el largo
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        lineas = max(1, -(-(len(etiqueta) + len(str(valor))) // CRECE[coord]))
        fila = ws.row_dimensions[c.row]
        fila.height = max(fila.height or 15, 12.5 * lineas + 3)
    if sep == "\n":
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        fila = ws.row_dimensions[c.row]
        fila.height = max(fila.height or 15, 24)


def llenar_htt(datos, salida):
    wb = load_workbook(PLANTILLAS / "htt.xlsx")
    ws = wb.active
    f = datos["fecha"]
    hh, _, mm = datos.get("hora", "").partition(":")
    v = {
        "dia": f"{f.day:02d}", "mes": f"{f.month:02d}", "anio": f.year, "hh": hh, "mm": mm,
        "ruta_eco": f'{datos.get("ruta", "")} / {datos.get("economico", "")}',
        "causal": datos.get("causal", ""), "colision": datos.get("colision", ""),
        "medio_prueba": datos.get("medio_prueba", ""),
        "reporte_aseguradora": datos.get("reporte_aseguradora", ""),
        "lugar": datos.get("lugar", ""), "colonia": datos.get("colonia", ""),
        "nota": datos.get("nota", ""),
        "ejecutivo_nombre": datos["a"].get("nombre", ""),
        "agente_nombre": datos.get("agente_nombre", ""),
        "perito_nombre": datos.get("perito_nombre", ""),
    }
    for k in ("a", "b"):
        for campo, valor in datos[k].items():
            v[f"{k}_{campo}"] = valor

    if "W7:AE7" not in [str(r) for r in ws.merged_cells.ranges]:
        ws.merge_cells("W7:AE7")  # para que "ruta / económico" no se corte
    for k, coord in VALOR.items():
        if v.get(k) not in (None, ""):
            _escribir(ws, coord, v[k])
    for k, coord in ETIQUETA.items():
        if v.get(k) not in (None, ""):
            _con_etiqueta(ws, coord, v[k])

    _marcar(ws, DIAS_SEMANA[f.weekday()])
    if datos.get("competencia") in COMPETENCIA:
        _marcar(ws, COMPETENCIA[datos["competencia"]])
    if datos.get("competencia") == "OTROS" and datos.get("competencia_otro"):
        _escribir(ws, "W9", datos["competencia_otro"])
    if datos.get("responsable") is not None:
        _marcar(ws, RESPONSABLE[datos["responsable"]])

    for k in ("a", "b"):
        p = datos[k]
        if p.get("sexo") in ("M", "F"):
            _marcar(ws, SEXO[k][p["sexo"]])
        if p.get("nacimiento"):  # date
            digitos = p["nacimiento"].strftime("%d%m%y")
            for coord, dgt in zip(NACIMIENTO[k], digitos):
                _escribir(ws, coord, dgt, sz=8)
        if p.get("lesionado"):
            _marcar(ws, LESIONADO[k])
        if p.get("fallecido"):
            _marcar(ws, FALLECIDO[k])

    for fila, les in zip(LESIONADOS_FILAS, datos.get("lesionados", [])):
        for campo, col in LESIONADOS_COLS:
            _escribir(ws, f"{col}{fila}", les.get(campo, ""))

    # anexos: lista de 6 dicts {"aplica": True/False/None, "art": "", "fracc": ""}
    for fila, anexo in zip(ANEXOS, datos.get("anexos", [])):
        if anexo.get("aplica") is True:
            _marcar(ws, f"A{fila}")
        elif anexo.get("aplica") is False:
            _marcar(ws, f"B{fila}")
        if anexo.get("art") or anexo.get("fracc"):
            valores = iter([anexo.get("art", ""), anexo.get("fracc", "")])
            ws[f"C{fila}"] = re.sub(r"_{3,}", lambda m: f" {next(valores, '')} ", ws[f"C{fila}"].value, count=2)

    for k, coord in FIRMAS.items():
        if datos.get(k):
            img = XLImage(datos[k])
            img.width, img.height = 130, 38  # px, cabe en el renglón de firma
            ws.add_image(img, coord)

    ruta = Path(salida) / "hecho_transito.xlsx"
    wb.save(ruta)
    return ruta


def generar_todo(datos, salida):
    Path(salida).mkdir(parents=True, exist_ok=True)
    archivos = [llenar_entrevista(datos, salida), llenar_htt(datos, salida)]
    if datos.get("aplica_convenio"):
        archivos.append(llenar_convenio(datos, salida))
    return archivos
